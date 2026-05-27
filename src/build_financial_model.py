"""Generate an interactive Excel financial model for Dopamine fasting and Picky eater.

Workbook structure:
  - Assumptions (editable inputs, three scenarios per product)
  - Dopamine — 12-month projection (formulas linked to Assumptions)
  - Picky — 12-month projection (formulas linked to Assumptions)
  - Summary (compares both at year-end)

User can edit Assumptions and all formulas recompute.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

OUT = Path(__file__).resolve().parent.parent / "output" / "financial_model.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)


BOLD = Font(bold=True, color="000000")
HEADER = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION = Font(bold=True, size=12, color="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # editable cells (light yellow)
CALC_FILL = PatternFill("solid", fgColor="E7E6E6")   # derived cells (light gray)
MONEY = "$#,##0"
PERCENT = "0.0%"
INT = "#,##0"
THIN = Side(border_style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_header(ws, row, cells):
    for col_idx, value in enumerate(cells, start=1):
        c = ws.cell(row=row, column=col_idx, value=value)
        c.font = HEADER
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER


def write_section(ws, row, col, title):
    c = ws.cell(row=row, column=col, value=title)
    c.font = SECTION


def write_input(ws, row, col, value, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = INPUT_FILL
    if fmt:
        c.number_format = fmt
    c.border = BORDER
    return c


def write_calc(ws, row, col, formula, fmt=None):
    c = ws.cell(row=row, column=col, value=formula)
    c.fill = CALC_FILL
    if fmt:
        c.number_format = fmt
    c.border = BORDER
    return c


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    ws.column_dimensions["A"].width = 40
    for col in "BCDE":
        ws.column_dimensions[col].width = 16

    ws["A1"] = "Financial Model — Assumptions"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Edit yellow cells. Three scenarios per product."
    ws["A2"].font = Font(italic=True, color="666666")

    # ── DOPAMINE FASTING ──
    write_section(ws, 4, 1, "Dopamine fasting")
    write_header(ws, 5, ["Parameter", "Pessimistic", "Realistic", "Optimistic"])

    dopamine_rows = [
        ("Month 1 installs",                          200,      500,      1000,    INT),
        ("Monthly install growth rate",               0.20,     0.40,     0.60,    PERCENT),
        ("Install → trial start (after Family Controls auth)", 0.10, 0.15,  0.20, PERCENT),
        ("Trial → paid conversion",                   0.20,     0.30,     0.40,    PERCENT),
        ("Annual subscriber share (vs monthly)",      0.50,     0.60,     0.70,    PERCENT),
        ("Monthly subscription price ($)",            5.99,     5.99,     5.99,    MONEY),
        ("Annual subscription price ($)",             39.99,    39.99,    39.99,   MONEY),
        ("Monthly churn (paying users)",              0.10,     0.06,     0.04,    PERCENT),
        ("Blended CAC per paid user ($)",             40,       15,       5,       MONEY),
    ]
    start_row = 6
    for i, (label, pess, real, opt, fmt) in enumerate(dopamine_rows):
        ws.cell(row=start_row + i, column=1, value=label).font = BOLD
        for col, v in [(2, pess), (3, real), (4, opt)]:
            write_input(ws, start_row + i, col, v, fmt)

    # ── PICKY EATER ──
    p_start = start_row + len(dopamine_rows) + 3
    write_section(ws, p_start - 1, 1, "Picky eater toddler")
    write_header(ws, p_start, ["Parameter", "Pessimistic", "Realistic", "Optimistic"])

    picky_rows = [
        ("Month 1 installs",                          100,      300,      600,     INT),
        ("Monthly install growth rate",               0.20,     0.35,     0.50,    PERCENT),
        ("Install → trial start",                     0.20,     0.25,     0.30,    PERCENT),
        ("Trial → paid conversion",                   0.20,     0.25,     0.35,    PERCENT),
        ("Monthly sub share",                         0.50,     0.50,     0.40,    PERCENT),
        ("Annual sub share",                          0.35,     0.35,     0.40,    PERCENT),
        ("Family sub share",                          0.15,     0.15,     0.20,    PERCENT),
        ("Monthly subscription ($)",                  9.99,     9.99,     9.99,    MONEY),
        ("Annual subscription ($/yr)",                59,       59,       59,      MONEY),
        ("Family monthly ($)",                        14.99,    14.99,    14.99,   MONEY),
        ("Monthly churn (paying)",                    0.15,     0.10,     0.07,    PERCENT),
        ("Blended CAC per paid user ($)",             100,      45,       20,      MONEY),
        ("Fixed monthly costs — RDN partner ($)",     1500,     1500,     1500,    MONEY),
        ("Fixed monthly costs — infrastructure ($)",  300,      300,      300,     MONEY),
    ]
    for i, (label, pess, real, opt, fmt) in enumerate(picky_rows):
        ws.cell(row=p_start + 1 + i, column=1, value=label).font = BOLD
        for col, v in [(2, pess), (3, real), (4, opt)]:
            write_input(ws, p_start + 1 + i, col, v, fmt)

    # legend
    legend_row = p_start + len(picky_rows) + 4
    ws.cell(row=legend_row, column=1, value="Legend:").font = BOLD
    c = ws.cell(row=legend_row + 1, column=1, value="Yellow — input (edit me)")
    c.fill = INPUT_FILL
    c2 = ws.cell(row=legend_row + 2, column=1, value="Gray — calculated (don't edit)")
    c2.fill = CALC_FILL


def build_dopamine_sheet(wb, scenario_col_letter="C"):
    """scenario_col_letter is the column in Assumptions for the chosen scenario (default Realistic=C)."""
    ws = wb.create_sheet("Dopamine — Realistic")
    ws.column_dimensions["A"].width = 32
    for col in "BCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 12

    ws["A1"] = "Dopamine fasting — 12-month projection (Realistic scenario)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Change scenario by editing formulas (default reads from Assumptions col C). Inputs in Assumptions sheet."
    ws["A2"].font = Font(italic=True, color="666666")

    # Header row
    months = ["Metric"] + [f"M{i}" for i in range(1, 13)] + ["Year"]
    write_header(ws, 4, months)

    s = scenario_col_letter  # column in Assumptions sheet

    # ROW 5: Monthly installs
    ws.cell(row=5, column=1, value="Monthly installs").font = BOLD
    # M1 = Assumptions!Cn (month-1 installs)
    write_calc(ws, 5, 2, f"=Assumptions!{s}6", INT)
    for m in range(2, 13):
        # subsequent month = previous * (1 + growth rate)
        prev = get_column_letter(m)  # column B for M1, etc.
        write_calc(ws, 5, m + 1, f"={prev}5*(1+Assumptions!{s}7)", INT)
    # Year total
    write_calc(ws, 5, 14, "=SUM(B5:M5)", INT)

    # ROW 6: Cumulative installs
    ws.cell(row=6, column=1, value="Cumulative installs").font = BOLD
    write_calc(ws, 6, 2, "=B5", INT)
    for m in range(2, 13):
        prev = get_column_letter(m)
        cur = get_column_letter(m + 1)
        write_calc(ws, 6, m + 1, f"={prev}6+{cur}5", INT)
    write_calc(ws, 6, 14, "=M6", INT)

    # ROW 7: Trials started this month
    ws.cell(row=7, column=1, value="Trials started").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 7, m + 1, f"={col}5*Assumptions!{s}8", INT)
    write_calc(ws, 7, 14, "=SUM(B7:M7)", INT)

    # ROW 8: New paying this month (trials × conversion)
    ws.cell(row=8, column=1, value="New paying this month").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 8, m + 1, f"={col}7*Assumptions!{s}9", INT)
    write_calc(ws, 8, 14, "=SUM(B8:M8)", INT)

    # ROW 9: Cumulative paying (with churn)
    ws.cell(row=9, column=1, value="Active paying users (after churn)").font = BOLD
    # M1: new paying this month only
    write_calc(ws, 9, 2, "=B8", INT)
    for m in range(2, 13):
        prev = get_column_letter(m)
        cur = get_column_letter(m + 1)
        # prev * (1 - churn) + new
        write_calc(ws, 9, m + 1, f"={prev}9*(1-Assumptions!{s}13)+{cur}8", INT)
    write_calc(ws, 9, 14, "=M9", INT)

    # ROW 10: ARPU per month (weighted blend monthly+annual)
    ws.cell(row=10, column=1, value="ARPU per paying user / mo ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        # = annual_share * (annual_price/12) + monthly_share * monthly_price
        write_calc(ws, 10, m + 1,
                   f"=Assumptions!{s}10*(Assumptions!{s}12/12)+(1-Assumptions!{s}10)*Assumptions!{s}11",
                   "$#,##0.00")
    write_calc(ws, 10, 14, "=M10", "$#,##0.00")

    # ROW 11: MRR
    ws.cell(row=11, column=1, value="MRR ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 11, m + 1, f"={col}9*{col}10", MONEY)
    write_calc(ws, 11, 14, "=M11", MONEY)

    # ROW 12: Monthly revenue (= MRR for simplicity since we're using mo ARPU)
    ws.cell(row=12, column=1, value="Monthly revenue ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 12, m + 1, f"={col}11", MONEY)
    write_calc(ws, 12, 14, "=SUM(B12:M12)", MONEY)

    # ROW 13: CAC spend this month
    ws.cell(row=13, column=1, value="CAC spend ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 13, m + 1, f"={col}8*Assumptions!{s}14", MONEY)
    write_calc(ws, 13, 14, "=SUM(B13:M13)", MONEY)

    # ROW 14: Net cash flow this month
    ws.cell(row=14, column=1, value="Net cash flow ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 14, m + 1, f"={col}12-{col}13", MONEY)
    write_calc(ws, 14, 14, "=SUM(B14:M14)", MONEY)

    # ROW 15: Cumulative cash
    ws.cell(row=15, column=1, value="Cumulative cash ($)").font = BOLD
    write_calc(ws, 15, 2, "=B14", MONEY)
    for m in range(2, 13):
        prev = get_column_letter(m)
        cur = get_column_letter(m + 1)
        write_calc(ws, 15, m + 1, f"={prev}15+{cur}14", MONEY)
    write_calc(ws, 15, 14, "=M15", MONEY)

    # ── Sensitivity table ──
    ws["A18"] = "Year-end snapshot"
    ws["A18"].font = SECTION
    ws["A19"] = "Total downloads"
    write_calc(ws, 19, 2, "=N6", INT)
    ws["A20"] = "Active paying users"
    write_calc(ws, 20, 2, "=M9", INT)
    ws["A21"] = "MRR at month 12"
    write_calc(ws, 21, 2, "=M11", MONEY)
    ws["A22"] = "ARR (MRR × 12)"
    write_calc(ws, 22, 2, "=B21*12", MONEY)
    ws["A23"] = "Total revenue year 1"
    write_calc(ws, 23, 2, "=N12", MONEY)
    ws["A24"] = "Total CAC spend"
    write_calc(ws, 24, 2, "=N13", MONEY)
    ws["A25"] = "Net cash year 1"
    write_calc(ws, 25, 2, "=N14", MONEY)


def build_picky_sheet(wb, scenario_col_letter="C"):
    """For Picky, params row mapping in Assumptions starts at row 19 (after Dopamine block + 3 rows)."""
    # Determine picky start row in Assumptions:
    # Assumptions: row 1-2 header, row 4 section, row 5 header, rows 6-14 dopamine (9 rows),
    # row 17 section, row 18 header, rows 19-32 picky (14 rows)
    # So picky params start at row 19 in Assumptions
    p = 19  # picky base row in Assumptions

    ws = wb.create_sheet("Picky — Realistic")
    ws.column_dimensions["A"].width = 32
    for col in "BCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 12

    ws["A1"] = "Picky eater — 12-month projection (Realistic scenario)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Change scenario by editing formulas (default reads from Assumptions col C). Inputs in Assumptions sheet."
    ws["A2"].font = Font(italic=True, color="666666")

    months = ["Metric"] + [f"M{i}" for i in range(1, 13)] + ["Year"]
    write_header(ws, 4, months)

    s = scenario_col_letter

    # ROW 5: Installs
    ws.cell(row=5, column=1, value="Monthly installs").font = BOLD
    write_calc(ws, 5, 2, f"=Assumptions!{s}{p}", INT)
    for m in range(2, 13):
        prev = get_column_letter(m)
        write_calc(ws, 5, m + 1, f"={prev}5*(1+Assumptions!{s}{p+1})", INT)
    write_calc(ws, 5, 14, "=SUM(B5:M5)", INT)

    # ROW 6: Cumulative installs
    ws.cell(row=6, column=1, value="Cumulative installs").font = BOLD
    write_calc(ws, 6, 2, "=B5", INT)
    for m in range(2, 13):
        prev = get_column_letter(m); cur = get_column_letter(m + 1)
        write_calc(ws, 6, m + 1, f"={prev}6+{cur}5", INT)
    write_calc(ws, 6, 14, "=M6", INT)

    # ROW 7: Trials started
    ws.cell(row=7, column=1, value="Trials started").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 7, m + 1, f"={col}5*Assumptions!{s}{p+2}", INT)
    write_calc(ws, 7, 14, "=SUM(B7:M7)", INT)

    # ROW 8: New paying this month
    ws.cell(row=8, column=1, value="New paying this month").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 8, m + 1, f"={col}7*Assumptions!{s}{p+3}", INT)
    write_calc(ws, 8, 14, "=SUM(B8:M8)", INT)

    # ROW 9: Active paying (with churn)
    ws.cell(row=9, column=1, value="Active paying users (after churn)").font = BOLD
    write_calc(ws, 9, 2, "=B8", INT)
    for m in range(2, 13):
        prev = get_column_letter(m); cur = get_column_letter(m + 1)
        write_calc(ws, 9, m + 1, f"={prev}9*(1-Assumptions!{s}{p+10})+{cur}8", INT)
    write_calc(ws, 9, 14, "=M9", INT)

    # ROW 10: Blended ARPU per month
    # = monthly_share * monthly_price + annual_share * (annual_price/12) + family_share * family_price
    ws.cell(row=10, column=1, value="ARPU per paying user / mo ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(
            ws, 10, m + 1,
            f"=Assumptions!{s}{p+4}*Assumptions!{s}{p+7}"
            f"+Assumptions!{s}{p+5}*(Assumptions!{s}{p+8}/12)"
            f"+Assumptions!{s}{p+6}*Assumptions!{s}{p+9}",
            "$#,##0.00",
        )
    write_calc(ws, 10, 14, "=M10", "$#,##0.00")

    # ROW 11: MRR
    ws.cell(row=11, column=1, value="MRR ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 11, m + 1, f"={col}9*{col}10", MONEY)
    write_calc(ws, 11, 14, "=M11", MONEY)

    # ROW 12: Monthly revenue
    ws.cell(row=12, column=1, value="Monthly revenue ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 12, m + 1, f"={col}11", MONEY)
    write_calc(ws, 12, 14, "=SUM(B12:M12)", MONEY)

    # ROW 13: CAC spend
    ws.cell(row=13, column=1, value="CAC spend ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 13, m + 1, f"={col}8*Assumptions!{s}{p+11}", MONEY)
    write_calc(ws, 13, 14, "=SUM(B13:M13)", MONEY)

    # ROW 14: Fixed costs
    ws.cell(row=14, column=1, value="Fixed monthly costs ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 14, m + 1,
                   f"=Assumptions!{s}{p+12}+Assumptions!{s}{p+13}", MONEY)
    write_calc(ws, 14, 14, "=SUM(B14:M14)", MONEY)

    # ROW 15: Net cash flow
    ws.cell(row=15, column=1, value="Net cash flow ($)").font = BOLD
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        write_calc(ws, 15, m + 1, f"={col}12-{col}13-{col}14", MONEY)
    write_calc(ws, 15, 14, "=SUM(B15:M15)", MONEY)

    # ROW 16: Cumulative cash
    ws.cell(row=16, column=1, value="Cumulative cash ($)").font = BOLD
    write_calc(ws, 16, 2, "=B15", MONEY)
    for m in range(2, 13):
        prev = get_column_letter(m); cur = get_column_letter(m + 1)
        write_calc(ws, 16, m + 1, f"={prev}16+{cur}15", MONEY)
    write_calc(ws, 16, 14, "=M16", MONEY)

    # Year-end snapshot
    ws["A19"] = "Year-end snapshot"
    ws["A19"].font = SECTION
    ws["A20"] = "Total downloads"; write_calc(ws, 20, 2, "=N6", INT)
    ws["A21"] = "Active paying users"; write_calc(ws, 21, 2, "=M9", INT)
    ws["A22"] = "MRR at month 12"; write_calc(ws, 22, 2, "=M11", MONEY)
    ws["A23"] = "ARR (MRR × 12)"; write_calc(ws, 23, 2, "=B22*12", MONEY)
    ws["A24"] = "Total revenue year 1"; write_calc(ws, 24, 2, "=N12", MONEY)
    ws["A25"] = "Total CAC spend"; write_calc(ws, 25, 2, "=N13", MONEY)
    ws["A26"] = "Total fixed costs"; write_calc(ws, 26, 2, "=N14", MONEY)
    ws["A27"] = "Net cash year 1"; write_calc(ws, 27, 2, "=N15", MONEY)


def build_summary(wb):
    ws = wb.create_sheet("Summary", 0)  # first sheet
    ws.column_dimensions["A"].width = 32
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    ws["A1"] = "Финансовая модель — сводка"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Сравнение сценариев. Поменяйте input в Assumptions — всё пересчитается."
    ws["A2"].font = Font(italic=True, color="666666")

    write_header(ws, 4, ["Метрика", "Dopamine (Realistic)", "Picky (Realistic)", "Разница"])

    rows = [
        ("Downloads year 1",    "='Dopamine — Realistic'!N6",  "='Picky — Realistic'!N6",  INT),
        ("Paying users at M12", "='Dopamine — Realistic'!M9",  "='Picky — Realistic'!M9",  INT),
        ("MRR at M12",           "='Dopamine — Realistic'!M11", "='Picky — Realistic'!M11", MONEY),
        ("ARR (×12)",            "='Dopamine — Realistic'!M11*12", "='Picky — Realistic'!M11*12", MONEY),
        ("Revenue year 1",       "='Dopamine — Realistic'!N12", "='Picky — Realistic'!N12", MONEY),
        ("CAC spend year 1",     "='Dopamine — Realistic'!N13", "='Picky — Realistic'!N13", MONEY),
        ("Net cash year 1",      "='Dopamine — Realistic'!N14", "='Picky — Realistic'!N15", MONEY),
    ]
    for i, (label, a, b, fmt) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = BOLD
        write_calc(ws, r, 2, a, fmt)
        write_calc(ws, r, 3, b, fmt)
        write_calc(ws, r, 4, f"=B{r}-C{r}", fmt)

    # Verdict
    ws["A14"] = "Что говорит модель"
    ws["A14"].font = SECTION
    ws["A15"] = ("Dopamine — solo-friendly: низкие издержки, высокий LTV/CAC, но "
                 "потолок ниже из-за более узкой аудитории и низкой цены.")
    ws["A15"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[15].height = 30
    ws["A16"] = ("Picky — выше потолок MRR и контент-moat, но требует RDN-партнёра "
                 "($1500/мес fixed) и партнёрства до запуска. Без partnership = риск.")
    ws["A16"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[16].height = 30
    ws["A18"] = "Откройте Assumptions, поменяйте жёлтые ячейки → пересчёт автоматический."
    ws["A18"].font = Font(italic=True, color="1F4E78", bold=True)


def main():
    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet
    build_assumptions(wb)
    build_dopamine_sheet(wb, "C")
    build_picky_sheet(wb, "C")
    build_summary(wb)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
