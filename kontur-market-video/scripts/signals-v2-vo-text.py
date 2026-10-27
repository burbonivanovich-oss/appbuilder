"""Текст диктора «ИИ Бизнес сигналы» v2 с таймкодами — отдельным файлом для записи.

  python3 scripts/signals-v2-vo-text.py   → Текст_диктора_ИИ_Бизнес_сигналы.docx

Таймкоды — из src/signals-v2/timeline.json, тексты — из src/signals-v2/script.json.
"""
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

T = json.load(open('src/signals-v2/timeline.json'))
S = json.load(open('src/signals-v2/script.json'))
SAY = {s['id']: s['say'] for c in S['chapters'] for s in c['scenes'] if 'say' in s}
OUT = 'Текст_диктора_ИИ_Бизнес_сигналы.docx'
GRAY = RGBColor(0x70, 0x70, 0x70)
BLUE = RGBColor(0x22, 0x91, 0xFF)


def tc(x):
    return f'{int(x // 60)}:{x % 60:04.1f}'.replace('.', ',')


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    pr.append(sh)


doc = Document()
st = doc.styles['Normal']
st.font.name = 'Arial'
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)


def para(text, size=11, bold=False, color=None, after=4, before=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


para('Текст диктора — «ИИ Бизнес сигналы»', 18, True, after=2)
para(f'Контур.Маркет · полная версия {tc(T["total"])[:-2]} · 12 реплик · время — мин:сек в полной версии, «в части» — от начала отдельного файла части', 10, color=GRAY, after=10)

para('Как записывать', 12, True, BLUE, after=4)
for t in [
    'Каждая реплика — отдельный файл WAV, 48 кГц / 24 бит, без музыки и обработки, с тишиной 0,5 с до и после. Имя файла — как в колонке «Файл».',
    'Темп — по длине в колонке «Файл · длина», допустимо ±10%: анимация подстроится под новую длину сама.',
    'Тон спокойный, уверенный, доброжелательный: деловой рассказ коллеге, не реклама.',
    '«ИИ» читается «И-И». «Контур.Маркет» — без паузы на точке. «2026 года» — «две тысячи двадцать шестого года».',
]:
    doc.add_paragraph(t, style='List Bullet').paragraph_format.space_after = Pt(2)

n = 0
for i, c in enumerate(T['chapters'], 1):
    para(f'Часть {i}. {c["title"]} · {tc(c["start"])}–{tc(c["start"] + c["dur"])}', 12, True, BLUE, before=12).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=5)
    t.style = 'Table Grid'
    for k, h in enumerate(['№', 'Время', 'В части', 'Файл · длина', 'Текст']):
        cell = t.rows[0].cells[k]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        shade(cell, 'E8F3FF')
    for s in c['scenes']:
        row = t.add_row().cells
        if 'vo' not in s:
            vals = ['—', f'{tc(s["start"])}–{tc(s["start"] + s["dur"])}', f'{tc(s["start"] - c["start"])}', '—', 'Заставка, без голоса']
        else:
            n += 1
            a = s['start'] + s['vo']['at']
            b = a + s['vo']['dur']
            vals = [str(n), f'{tc(a)}–{tc(b)}', f'{tc(a - c["start"])}–{tc(b - c["start"])}', f'{s["id"]}.wav\n' + f'{s["vo"]["dur"]:.1f} с'.replace('.', ','), SAY[s['id']]]
        for k, v in enumerate(vals):
            row[k].text = ''
            r = row[k].paragraphs[0].add_run(v)
            r.font.size = Pt(9 if k < 4 else 11)
            if k in (2, 3):
                r.font.color.rgb = GRAY
    # фиксированная раскладка: иначе Word/LibreOffice сжимают колонку с текстом
    t.autofit = False
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    t._tbl.tblPr.append(layout)
    widths = [0.8, 2.6, 2.6, 2.4, 9.0]
    for k, w in enumerate(widths):
        t.columns[k].width = Cm(w)
    for row in t.rows:
        for k, w in enumerate(widths):
            row.cells[k].width = Cm(w)

para('Финальный кадр отдельных частей (2,8 с, после главы) — без голоса.', 9, color=GRAY, before=10)
doc.save(OUT)
print(OUT)
